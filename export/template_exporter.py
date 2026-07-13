import json
import re
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.worksheet import Worksheet

from config import settings
from extraction.prompts import TEMPLATE_MAPPING_PROMPT
from models.exceptions import ModelQueryError
from models.router import query_vision_model, ModelBackend
from utils.helpers import logger


def _detect_file_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix in (".xlsx", ".xls"):
        return "xlsx"
    raise ValueError(f"Unsupported template format: {suffix}. Use .xlsx or .csv")


def read_template_headers(path: Path) -> list[str]:
    ext = _detect_file_type(path)
    if ext == "csv":
        df = pd.read_csv(path, nrows=0)
        return list(df.columns)
    df = pd.read_excel(path, nrows=0, engine="openpyxl")
    return list(df.columns)


def map_variables_to_columns(
    variable_names: list[str],
    template_headers: list[str],
    study_id: str,
) -> dict[str, str]:
    prompt = TEMPLATE_MAPPING_PROMPT.format(
        variable_names=json.dumps(variable_names),
        template_headers=json.dumps(template_headers),
    )
    try:
        response = query_vision_model(prompt)
        mapping = _parse_mapping_response(response, variable_names, template_headers)
        if any(mapping.values()):
            return mapping
    except ModelQueryError as e:
        logger.warning(f"LLM mapping failed ({e}); falling back to heuristic mapping")

    return _heuristic_mapping(variable_names, template_headers)


def _heuristic_mapping(variable_names: list[str], template_headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for var in variable_names:
        var_lower = var.lower().replace("_", "").replace(" ", "")
        best_match = ""
        best_score = 0
        for header in template_headers:
            header_lower = header.lower().replace("_", "").replace(" ", "")
            score = _similarity_score(var_lower, header_lower)
            if score > best_score:
                best_score = score
                best_match = header
        mapping[var] = best_match if best_score >= 0.3 else ""
    return mapping


def _similarity_score(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.8
    common = len(set(a) & set(b))
    total = len(set(a) | set(b))
    return common / total if total > 0 else 0.0


def _parse_mapping_response(
    response: str,
    variable_names: list[str],
    template_headers: list[str],
) -> dict[str, str]:
    json_match = re.search(r"\{.*\}", response, re.DOTALL)
    if not json_match:
        return {}
    try:
        parsed: dict[str, str | list | None] = json.loads(json_match.group(0))
        mapping_root = parsed.get("mapping", parsed)
        if isinstance(mapping_root, dict):
            parsed = mapping_root
        flat: dict[str, str] = {}
        for var in variable_names:
            raw = parsed.get(var, "")
            if isinstance(raw, list):
                matched = raw[0] if raw else ""
            elif raw is None:
                matched = ""
            else:
                matched = str(raw)
            if matched and matched not in template_headers:
                logger.warning(
                    f"Mapped column '{matched}' for '{var}' not in template headers"
                )
                matched = ""
            flat[var] = matched
        return flat
    except json.JSONDecodeError:
        logger.warning("Failed to parse mapping JSON response")
        return {}


def _collect_variable_names(result_dict: dict) -> list[str]:
    names: list[str] = []
    for arm in result_dict.get("arms", []):
        for cd in arm.get("continuous", []):
            v = cd.get("variable")
            if v and v not in names:
                names.append(v)
        for dd in arm.get("dichotomous", []):
            v = dd.get("variable")
            if v and v not in names:
                names.append(v)
    return names


def _build_flat_row(
    result_dict: dict, mapping: dict[str, str], study_id: str
) -> dict[str, Any]:
    row: dict[str, Any] = {}
    seen_columns: set[str] = set()

    for arm in result_dict.get("arms", []):
        arm_label = arm.get("label", "Unknown")
        for cd in arm.get("continuous", []):
            var_col = cd.get("variable", "")
            mapped = mapping.get(var_col, "")
            if mapped:
                col = _qualify_column(mapped, arm_label, seen_columns)
                seen_columns.add(col)
                row[col] = cd.get("mean")
                row[f"{col}_SD"] = cd.get("sd")
                row[f"{col}_N"] = cd.get("n")

        for dd in arm.get("dichotomous", []):
            var_col = dd.get("variable", "")
            mapped = mapping.get(var_col, "")
            if mapped:
                col = _qualify_column(mapped, arm_label, seen_columns)
                seen_columns.add(col)
                row[f"{col}_Events"] = dd.get("events")
                row[f"{col}_Total"] = dd.get("total")

    row["Study_ID"] = study_id
    return row


def _qualify_column(base_col: str, arm_label: str, seen: set[str]) -> str:
    candidate = f"{base_col}_{arm_label.lower().replace(' ', '_')}"
    return candidate if candidate in seen else base_col


def _apply_row_to_worksheet(
    ws: Worksheet,
    row_data: dict[str, Any],
) -> int:
    new_row = ws.max_row + 1
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        header = cell.value
        if header is None:
            continue
        header_str = str(header).strip()
        value = row_data.get(header_str)
        target = ws.cell(row=new_row, column=col_idx)
        if value is not None:
            target.value = value
        else:
            target.value = "Not Found"
            target.comment = Comment(
                f"No data found in source for column '{header_str}'.",
                "DataExtractor_Noora",
            )
    ws.cell(row=new_row, column=1).value = row_data.get("Study_ID", "unknown")
    return new_row


def idx_to_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def export_to_template(
    extraction_path: Path,
    template_path: Path,
    output_path: Path,
) -> Path:
    with open(extraction_path, encoding="utf-8") as f:
        result_dict = json.load(f)

    logger.info(f"Reading template: {template_path}")
    ext = _detect_file_type(template_path)
    study_id = result_dict.get("study_id", "unknown")

    if ext == "csv":
        df = pd.read_csv(template_path)
        mapping = map_variables_to_columns(
            _collect_variable_names(result_dict), list(df.columns), study_id
        )
        row = _build_flat_row(result_dict, mapping, study_id)
        new_row_dict = {c: row.get(c, None) for c in df.columns}
        df.loc[len(df)] = new_row_dict
        if output_path.suffix.lower() == ".csv":
            df.to_csv(output_path, index=False)
        else:
            df.to_excel(output_path, index=False, engine="openpyxl")
        logger.info(f"CSV template populated: {output_path}")
        return output_path

    wb = load_workbook(template_path)
    ws = wb.active
    if ws is None:
        raise ValueError("Template has no active worksheet.")

    headers = [str(ws.cell(row=1, column=c).value or f"Col_{c}").strip()
               for c in range(1, ws.max_column + 1)]

    mapping = map_variables_to_columns(
        _collect_variable_names(result_dict), headers, study_id
    )
    row_data = _build_flat_row(result_dict, mapping, study_id)
    _apply_row_to_worksheet(ws, row_data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"XLSX template populated: {output_path}")
    return output_path
